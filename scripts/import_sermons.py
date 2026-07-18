from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import textwrap
import csv
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "raw" / "教会讲道"
ORGANIZED_DIR = ROOT / "data" / "processed" / "整理后的讲道文章"
POSTS_DIR = ROOT / "src" / "content" / "posts"
REPORT_DIR = ROOT / "docs" / "内容整理报告"

SPEAKER_TAGS = {
    "Patrick": "Patrick",
    "Grayson": "Grayson",
    "Ryan": "Ryan",
    "Akira": "Akira",
    "Sundy": "Sundy",
    "Kota": "Kota",
    "Kaz": "Kaz",
}

BOOKS = (
    "创世记", "出埃及记", "利未记", "民数记", "申命记", "约书亚记", "士师记", "路得记",
    "撒母耳记上", "撒母耳记下", "列王纪上", "列王纪下", "历代志上", "历代志下", "以斯拉记",
    "尼希米记", "以斯帖记", "约伯记", "诗篇", "箴言", "传道书", "雅歌", "以赛亚书",
    "耶利米书", "耶利米哀歌", "以西结书", "但以理书", "何西阿书", "约珥书", "阿摩司书",
    "俄巴底亚书", "约拿书", "约拿", "弥迦书", "那鸿书", "哈巴谷书", "西番雅书", "哈该书",
    "撒迦利亚书", "玛拉基书", "马太福音", "马可福音", "路加福音", "约翰福音", "使徒行传",
    "罗马书", "哥林多前书", "哥林多后书", "加拉太书", "以弗所书", "腓立比书", "歌罗西书",
    "帖撒罗尼迦前书", "帖撒罗尼迦后书", "提摩太前书", "提摩太后书", "提多书", "腓利门书",
    "希伯来书", "雅各书", "彼得前书", "彼得后书", "约翰一书", "约翰二书", "约翰三书",
    "犹大书", "启示录",
)

CHINESE_NUMBERS = {
    "一": "1", "二": "2", "三": "3", "四": "4", "五": "5",
    "六": "6", "七": "7", "八": "8", "九": "9", "十": "10",
}

ENGLISH_BOOKS = {
    "genesis": "创世记", "gen": "创世记",
    "exodus": "出埃及记", "exo": "出埃及记", "exod": "出埃及记",
    "psalms": "诗篇", "psalm": "诗篇", "ps": "诗篇",
    "matthew": "马太福音", "matt": "马太福音", "mt": "马太福音",
    "mark": "马可福音", "mk": "马可福音",
    "luke": "路加福音", "lk": "路加福音",
    "john": "约翰福音", "jn": "约翰福音",
    "acts": "使徒行传", "act": "使徒行传",
    "romans": "罗马书", "roman": "罗马书", "rom": "罗马书",
    "1 corinthians": "哥林多前书", "1 cor": "哥林多前书", "1cor": "哥林多前书",
    "2 corinthians": "哥林多后书", "2 cor": "哥林多后书", "2cor": "哥林多后书",
    "galatians": "加拉太书", "gal": "加拉太书",
    "ephesians": "以弗所书", "eph": "以弗所书",
    "philippians": "腓立比书", "phil": "腓立比书",
    "colossians": "歌罗西书", "col": "歌罗西书",
    "hebrews": "希伯来书", "heb": "希伯来书",
    "james": "雅各书", "jas": "雅各书",
    "1 peter": "彼得前书", "1 pet": "彼得前书", "1pet": "彼得前书",
    "2 peter": "彼得后书", "2 pet": "彼得后书", "2pet": "彼得后书",
    "1 john": "约翰一书", "1 jn": "约翰一书", "1jn": "约翰一书",
    "2 john": "约翰二书", "2 jn": "约翰二书", "2jn": "约翰二书",
    "3 john": "约翰三书", "3 jn": "约翰三书", "3jn": "约翰三书",
    "revelation": "启示录", "rev": "启示录",
}

TEXT_REPLACEMENTS = {
    "[TF] ": "",
    "分かち合い": "分享",
    "・": "·",
    "创世纪": "创世记",
    "约伯纪": "约伯记",
    "钉S字J": "钉十字架",
    "人囗": "人口",
    "无庸置疑": "毋庸置疑",
    "想像": "想象",
    "彷佛": "仿佛",
    "于于神": "于神",
    "神。 神。": "神。",
    "整整一整夜": "整整一夜",
    "【雅各书1:9-18诱惑": "雅各书1:9-18诱惑",
    "【2-被诱惑": "2-被诱惑",
    "【3 -被拖走": "3 -被拖走",
    "【4 -罪": "4 -罪",
    "【5 -死": "5 -死",
    "怎样ネ能": "怎样才能",
    "第二歩": "第二步",
    "咬ー口": "咬一口",
    "探讨“合一的生活\"": "探讨“合一的生活”",
    "与神 和他的旨意": "与神和他的旨意",
    "一起 阅读": "一起阅读",
}


def clean_title(value: str) -> str:
    value = value.replace("：", ":").replace("；", ";").replace("，", ";")
    value = value.replace("–", "-").replace("—", "-").replace("_", ":")
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"\s*([:;,\-])\s*", r"\1", value)
    value = re.sub(r"(\d+:\d+-\d+):(\d+-\d+)", r"\1;\2", value)
    return value.strip(" _-:;，。")


def english_book_pattern() -> str:
    return "|".join(sorted((re.escape(book) for book in ENGLISH_BOOKS), key=len, reverse=True))


def scripture_from_english(normalized: str) -> tuple[str, str]:
    match = re.search(
        rf"(?P<book>{english_book_pattern()})\.?\s*(?P<chapter>\d+)\s*[:.]\s*(?P<verses>\d+(?:-\d+:\d+|-\d+)?(?:[;,]\d+(?:-\d+)?)?)",
        normalized,
        flags=re.IGNORECASE,
    )
    if not match:
        return "", normalized
    book_key = re.sub(r"\s+", " ", match.group("book").lower().replace(".", "")).strip()
    book = ENGLISH_BOOKS.get(book_key, "")
    if not book:
        return "", normalized
    scripture = f"{book} {match.group('chapter')}:{match.group('verses').replace(',', ';')}"
    summary = (normalized[:match.start()] + normalized[match.end():]).strip(" _-:;，。")
    summary = re.sub(r"^(?:系列:)?", "", summary).strip(" _-:;，。")
    return scripture, summary


def title_parts(title: str) -> tuple[str, str]:
    normalized = clean_title(title)
    normalized = re.sub(r"^(?:中文|福音\d+)", "", normalized).strip()
    books = "|".join(sorted((re.escape(book) for book in BOOKS), key=len, reverse=True))
    match = re.search(
        rf"(?P<book>{books})\s*(?:第(?:(?P<cn>[一二三四五六七八九十]+)|(?P<chapter_named>\d+))章|(?P<chapter>\d+)(?P<range>-\d+)?章?)"
        rf"(?:\s*[:;,]\s*(?P<verses>\d+(?:-\d+:\d+|-\d+)?(?:[;,]\d+(?:-\d+)?)?))?",
        normalized,
    )
    if not match:
        return scripture_from_english(normalized)

    book = "约拿书" if match.group("book") == "约拿" else match.group("book")
    chapter = (
        match.group("chapter")
        or match.group("chapter_named")
        or CHINESE_NUMBERS.get(match.group("cn"), match.group("cn") or "")
    )
    chapter_range = match.group("range") or ""
    verses = match.group("verses") or ""
    scripture = f"{book} {chapter}{chapter_range}"
    if verses:
        scripture += f":{verses.replace(',', ';')}"
    summary = (normalized[:match.start()] + normalized[match.end():]).strip(" _-:;，。")
    summary = re.sub(r"^(?:系列:)?", "", summary).strip(" _-:;，。")
    return scripture, summary


def scripture_from_body(body: str) -> str:
    sample = re.sub(r"\s+", " ", body[:4000])
    scripture, _summary = title_parts(sample)
    return scripture


def resolve_scripture(raw_title: str, source_file: Path, body: str) -> tuple[str, str]:
    candidates: dict[str, str] = {}
    folder_scripture, _summary = title_parts(raw_title)
    file_scripture, _file_summary = title_parts(source_file.stem)
    body_scripture = scripture_from_body(body)
    if folder_scripture:
        candidates["folder"] = folder_scripture
    if file_scripture:
        candidates["file"] = file_scripture
    if body_scripture:
        candidates["body"] = body_scripture

    unique = sorted(set(candidates.values()))
    if len(unique) > 1:
        details = "; ".join(f"{source}={scripture}" for source, scripture in candidates.items())
        raise SystemExit(f"Scripture conflict detected; please confirm metadata manually: {details}")
    if unique:
        confidence = "high" if len(candidates) >= 2 else "medium"
        return unique[0], confidence
    return "", "low"


def category_for(title: str) -> str:
    return "教会讲道"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5", "shift_jis"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_docx(path: Path) -> str:
    doc = Document(path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    return "\n\n".join(lines)


def normalize_body(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    for old, new in TEXT_REPLACEMENTS.items():
        text = text.replace(old, new)
    lines = [line.strip() for line in text.splitlines()]
    blocks: list[str] = []
    current: list[str] = []

    for line in lines:
        if not line:
            if current:
                blocks.append(" ".join(current).strip())
                current = []
            continue
        if re.match(r"^#{1,6}\s+", line) or re.match(r"^[-*]\s+", line):
            if current:
                blocks.append(" ".join(current).strip())
                current = []
            blocks.append(line)
        else:
            current.append(line)

    if current:
        blocks.append(" ".join(current).strip())

    return "\n\n".join(blocks).strip()


def yaml_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace('"', '\\"')


def read_manual_description(folder: Path, explicit: str | None = None) -> str:
    if explicit:
        return explicit.strip()

    metadata_path = folder / "metadata.json"
    if metadata_path.exists():
        try:
            data = json.loads(metadata_path.read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as exc:
            raise SystemExit(f"Invalid metadata.json: {metadata_path} ({exc})")
        value = data.get("description") if isinstance(data, dict) else None
        if isinstance(value, str) and value.strip():
            return value.strip()

    description_path = folder / "description.txt"
    if description_path.exists():
        return description_path.read_text(encoding="utf-8-sig").strip()

    return ""


def validate_manual_description(description: str, body: str) -> None:
    normalized = re.sub(r"\s+", " ", description).strip()
    placeholders = {"NEEDS_DESCRIPTION_REVIEW", "TODO", "TBD", "待补", "暂无"}
    if not normalized:
        raise SystemExit("Missing manual description. Add --description, metadata.json, or description.txt.")
    if normalized in placeholders:
        raise SystemExit("Description is a placeholder; add a manual summary before publishing.")
    if len(normalized) < 30:
        raise SystemExit("Description is too short; add a complete manual summary.")
    if not re.search(r"[。！？.!?]$", normalized):
        raise SystemExit("Description must be a complete sentence ending with punctuation.")
    body_head = re.sub(r"\s+", " ", body[:300]).strip()
    if body_head and normalized[:40] in body_head:
        raise SystemExit("Description appears copied from the body opening; write a manual summary instead.")


def sentence_aware_description(body: str, fallback: str) -> str:
    blocks = [block.strip() for block in re.split(r"\n\s*\n", body) if block.strip()]
    sentences: list[str] = []
    for block in blocks[:8]:
        text = re.sub(r"\s+", " ", block).strip()
        if not text or text in {"⸻", "---"}:
            continue
        if len(text) <= 30 and not re.search(r"[。！？.!?]$", text):
            continue
        parts = re.findall(r".+?[。！？.!?](?=\s|$)|.+$", text)
        for part in parts:
            sentence = part.strip()
            if len(sentence) >= 24:
                sentences.append(sentence)
            if len("".join(sentences)) >= 80:
                return "".join(sentences)
    return "".join(sentences) or fallback


def slugify(value: str) -> str:
    value = value.lower()
    value = re.sub(r"[\s_：:;；,，.。()（）\[\]【】/\\]+", "-", value)
    value = re.sub(r"-+", "-", value).strip("-")
    return value or "sermon"


def folder_meta(folder: Path) -> tuple[str, str, str]:
    name = folder.name
    match = re.match(r"^(\d{8})(.*)$", name)
    if not match:
        raise ValueError(f"目录名没有日期开头: {name}")

    raw_date, rest = match.groups()
    date = f"{raw_date[:4]}-{raw_date[4:6]}-{raw_date[6:8]}"
    speaker = ""
    title = rest.strip(" _-")

    if "_" in title:
        before, after = title.rsplit("_", 1)
        if after.strip():
            title = before.strip()
            speaker = after.strip()

    for key, value in SPEAKER_TAGS.items():
        if key.lower() in name.lower():
            speaker = value
            break

    title = clean_title(title or name)
    return date, title, speaker


def choose_source_file(folder: Path) -> Path | None:
    txt_files = sorted(
        (
            path
            for path in folder.glob("*.txt")
            if not path.name.lower().endswith(".extracted.txt")
        ),
        key=lambda p: p.name,
    )
    docx_files = sorted(folder.glob("*.docx"), key=lambda p: p.name)

    preferred_txt = [
        p for p in txt_files
        if any(token in p.name for token in ("中文", "中", "文本"))
        and not any(token in p.name.lower() for token in ("英文", "英", "日文", "jp"))
    ]
    if preferred_txt:
        return preferred_txt[0]

    neutral_txt = [
        p for p in txt_files
        if not any(token in p.name.lower() for token in ("英文", "英", "日文", "jp"))
    ]
    if neutral_txt:
        return neutral_txt[0]

    preferred_docx = [
        p for p in docx_files
        if any(token in p.name for token in ("中文", "中"))
        and not any(token in p.name.lower() for token in ("英文", "英", "jp"))
    ]
    if preferred_docx:
        return preferred_docx[0]

    neutral_docx = [
        p for p in docx_files
        if not any(token in p.name.lower() for token in ("英文", "英", "jp"))
    ]
    if neutral_docx:
        return neutral_docx[0]

    return None


def markdown_for(folder: Path, source_file: Path, description_override: str | None = None) -> tuple[str, str]:
    _source_date, raw_title, speaker = folder_meta(folder)
    # Website display date must be the publication date, not the source folder date.
    published_at = datetime.now(ZoneInfo("Asia/Tokyo"))
    date = published_at.strftime("%Y-%m-%d")
    _folder_scripture, summary = title_parts(raw_title)
    category = category_for(raw_title)
    if source_file.suffix.lower() == ".docx":
        body = read_docx(source_file)
    else:
        body = decode_text(source_file)
    body = normalize_body(body)
    scripture, scripture_confidence = resolve_scripture(raw_title, source_file, body)
    if scripture_confidence == "low":
        raise SystemExit("No scripture could be identified; please confirm metadata manually before publishing.")
    if scripture:
        title = f"{scripture}｜{summary or '经文讲解'}"
    else:
        title = summary or raw_title

    description = read_manual_description(folder, description_override)
    validate_manual_description(description, body)
    tags = ["讲道", category]
    if speaker:
        tags.append(speaker)

    frontmatter = textwrap.dedent(
        f"""\
        ---
        title: "{yaml_escape(title)}"
        description: "{yaml_escape(description)}"
        date: {date}
        publishedAt: {published_at.isoformat(timespec="seconds")}
        tags: [{", ".join(f'"{yaml_escape(tag)}"' for tag in tags)}]
        category: "{yaml_escape(category)}"
        scripture: "{yaml_escape(scripture)}"
        scriptureConfidence: "{yaml_escape(scripture_confidence)}"
        author: "{yaml_escape(speaker)}"
        reviewed: false
        source: "data/raw/教会讲道/{yaml_escape(folder.name)}/{yaml_escape(source_file.name)}"
        ---
        """
    )
    return date, frontmatter + "\n" + body + "\n"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import one sermon folder safely.")
    parser.add_argument("--folder", help="Single sermon folder to import.")
    parser.add_argument("--dry-run", action="store_true", help="Preview without writing files.")
    parser.add_argument("--description", help="Manual frontmatter summary. Do not pass body excerpts or templates.")
    parser.add_argument("--update-existing", action="store_true", help="Update the already-registered post for this same source folder.")
    return parser.parse_args()


def resolve_folder(value: str) -> Path:
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        direct = (ROOT / candidate).resolve()
        if direct.exists():
            candidate = direct
        else:
            candidate = (SOURCE_DIR / value).resolve()
    else:
        candidate = candidate.resolve()
    try:
        candidate.relative_to(SOURCE_DIR.resolve())
    except ValueError:
        raise SystemExit(f"Sermon folder must be under {SOURCE_DIR}: {candidate}")
    if not candidate.is_dir():
        raise SystemExit(f"Sermon folder not found: {candidate}")
    return candidate


def target_paths_for(source_date: str, publish_title: str) -> tuple[str, Path, Path]:
    slug = f"{source_date}-{slugify(publish_title)}"
    return slug, ORGANIZED_DIR / f"{slug}.md", POSTS_DIR / f"{slug}.md"


def read_sermon_registry(path: Path) -> dict[str, dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            source_folder = row.get("source_folder", "").strip()
            if source_folder:
                rows[source_folder] = dict(row)
    return rows


def write_sermon_registry(path: Path, rows: dict[str, dict[str, str]]) -> None:
    fieldnames = (
        "source_folder",
        "source_file",
        "source_sha256",
        "scripture",
        "scripture_confidence",
        "slug",
        "processed_path",
        "post_path",
        "first_imported_date",
        "last_checked_date",
        "status",
    )
    path.parent.mkdir(exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in sorted(rows.values(), key=lambda item: item.get("source_folder", "")):
            writer.writerow({name: row.get(name, "") for name in fieldnames})


def ensure_safe_targets(organized_path: Path, post_path: Path, *, update_existing: bool = False) -> None:
    existing = [path for path in (organized_path, post_path) if path.exists()]
    if existing and not update_existing:
        details = "\n".join(f"- {path}" for path in existing)
        raise SystemExit(
            "Target file already exists; stopping to avoid overwriting existing posts.\n"
            f"{details}\n"
            "Use --update-existing only when updating the same registered source folder."
        )
    missing = [path for path in (organized_path, post_path) if not path.exists()]
    if update_existing and missing:
        details = "\n".join(f"- {path}" for path in missing)
        raise SystemExit(
            "Registered update target is missing; stopping instead of creating a partial duplicate.\n"
            f"{details}"
        )


def main() -> None:
    args = parse_args()
    if not args.folder:
        raise SystemExit(
            "import_sermons.py no longer imports all sermon folders by default. "
            "Use --folder to select one folder; add --dry-run to preview."
        )

    if not SOURCE_DIR.exists():
        raise SystemExit(f"找不到源目录: {SOURCE_DIR}")

    target_folder = resolve_folder(args.folder)

    if not args.dry_run:
        ORGANIZED_DIR.mkdir(exist_ok=True)
        POSTS_DIR.mkdir(parents=True, exist_ok=True)
        REPORT_DIR.mkdir(exist_ok=True)

    report_path = REPORT_DIR / "讲道文章目录.csv"
    registry_path = REPORT_DIR / "sermon-import-registry.csv"
    registry_rows = read_sermon_registry(registry_path)
    rows_by_folder: dict[str, tuple[str, ...]] = {}
    if report_path.exists():
        with report_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            next(reader, None)
            for row in reader:
                if len(row) == 7:
                    rows_by_folder[row[0]] = tuple(row)

    imported = 0
    skipped: list[str] = []
    for folder in sorted(p for p in SOURCE_DIR.iterdir() if p.is_dir()):
        if folder.resolve() != target_folder:
            continue
        source_file = choose_source_file(folder)
        if source_file is None:
            skipped.append(folder.name)
            continue

        source_hash = file_sha256(source_file)
        date, markdown = markdown_for(folder, source_file, args.description)
        source_date, _raw_title, _speaker = folder_meta(folder)
        title_match = re.search(r'^title: "([^"]+)"', markdown, flags=re.MULTILINE)
        publish_title = title_match.group(1) if title_match else folder.name
        existing_registry = registry_rows.get(folder.name)
        if existing_registry and not args.update_existing:
            raise SystemExit(
                "Source folder already exists in sermon-import-registry.csv; "
                f"stopping to avoid duplicate publish: {folder.name} -> {existing_registry.get('slug', '')}; "
                f"old sha256={existing_registry.get('source_sha256', '')}; new sha256={source_hash}"
            )
        if existing_registry:
            slug = existing_registry.get("slug", "")
            processed_value = existing_registry.get("processed_path", "")
            post_value = existing_registry.get("post_path", "")
            if not slug or not processed_value or not post_value:
                raise SystemExit("Existing registry row is incomplete; cannot update safely.")
            organized_path = ROOT / processed_value
            post_path = ROOT / post_value
        else:
            slug, organized_path, post_path = target_paths_for(source_date, publish_title)
        if any(row.get("slug") == slug and row.get("source_folder") != folder.name for row in registry_rows.values()):
            raise SystemExit(f"Slug already belongs to another source folder: {slug}")
        ensure_safe_targets(organized_path, post_path, update_existing=bool(existing_registry and args.update_existing))
        print(f"Folder: {folder.name}")
        print(f"Source file: {source_file.name}")
        print(f"Title: {publish_title}")
        print(f"Slug: {slug}")
        print(f"Processed target: {organized_path}")
        print(f"Post target: {post_path}")
        if args.dry_run:
            print("Dry-run: no files written.")
        else:
            organized_path.write_text(markdown, encoding="utf-8")
            shutil.copyfile(organized_path, post_path)

        _source_date, raw_title, speaker = folder_meta(folder)
        scripture, _summary = title_parts(raw_title)
        if not scripture:
            scripture_match = re.search(r'^scripture: "([^"]+)"', markdown, flags=re.MULTILINE)
            scripture = scripture_match.group(1) if scripture_match else ""
        confidence_match = re.search(r'^scriptureConfidence: "([^"]+)"', markdown, flags=re.MULTILINE)
        scripture_confidence = confidence_match.group(1) if confidence_match else ""
        category = category_for(raw_title)
        rows_by_folder[folder.name] = (
            folder.name, date, scripture, category, speaker, publish_title, source_file.name
        )
        registry_rows[folder.name] = {
            "source_folder": folder.name,
            "source_file": source_file.name,
            "source_sha256": source_hash,
            "scripture": scripture,
            "scripture_confidence": scripture_confidence,
            "slug": slug,
            "processed_path": str(organized_path.relative_to(ROOT)),
            "post_path": str(post_path.relative_to(ROOT)),
            "first_imported_date": existing_registry.get("first_imported_date", date) if existing_registry else date,
            "last_checked_date": date,
            "status": "updated" if existing_registry else "imported",
        }
        imported += 1

    if not args.dry_run:
        with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(("源目录", "日期", "经文", "分类", "讲员", "发布标题", "正文来源"))
            writer.writerows(sorted(rows_by_folder.values(), key=lambda row: (row[1], row[0])))
        write_sermon_registry(registry_path, registry_rows)

    label = "Would import" if args.dry_run else "Imported"
    print(f"{label}: {imported}")
    print(f"Skipped: {len(skipped)}")
    for name in skipped:
        print("Skipped: " + name.encode("unicode_escape").decode("ascii"))


if __name__ == "__main__":
    main()
